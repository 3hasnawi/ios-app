"""Builds a test payroll workbook with the same layout as the real monthly file,
filled with made-up names, for the simulator test. Real payroll files never go to the repo.

The payroll sheet keeps its in-sheet deduction columns (اسم الاستقطاع / مبلغ الاستقطاع) filled in,
so the test also checks that the app takes deductions from the deductions sheet only.

Usage: python make_test_xlsx.py OUT.xlsx EXPECTED.txt [EMPLOYEE_COUNT] [VARIANT]
VARIANT "previous" changes the fixed employee's salary and deductions, to tell two months apart.
"""
import random
import sys

from openpyxl import Workbook

out_path, expected_path = sys.argv[1], sys.argv[2]
count = int(sys.argv[3]) if len(sys.argv) > 3 else 26000
previous = len(sys.argv) > 4 and sys.argv[4] == "previous"
rng = random.Random(2025 if previous else 2026)

HEADERS = [
    "رقم الموظف", "الدرجة الوظيفية", "الرتبة", "Name_m", "الاسمي", "الخطورة", "الارزاق", "الاستثنائية",
    "الاطفال", "الزوجية", "الشهادة", "منصب", "مخصصات صنف", "نسبة الصنف", "اسم الصنف", "جامعية",
    "مخصصات نقل", "مؤقتة", "اضافية", "مخصصات اخرى", "لقب علمي", "اسم المخصصات الاخرى", "نسبة من الاسمي",
    "الاستحقاق", "توقيفات تقاعدية", "الرعاية الصحية", "صندوق الشهداء", "نشاط صحي", "نشاط رياضي", "الضريبة",
    "تضمين", "ايراد نهائي", "تامين صحي", "مجموع استقطاعات المصارف", "حماية الاجتماعية", "رسم طابع",
    "اسم الاستقطاع", "مبلغ الاستقطاع", "2اسم الاستقطاع", "2مبلغ الاستقطاع", "% استقطاع 3",
    "مجموع الاستقطاعات", "صافي الراتب", "ملاحظات عامه", "NOTA1", "ملاحظات", "ملاحظات الغياب",
    "ماحظات التنقلات", "ملاحظات التسوية", "اسم الوكالة", "اسم المديرية", "NO_markaz", "رقم الضمان الصحي",
    "مدخل البانات واسم الحاسبة ", "تاريخ التعديل", "اسم الشهدة", "اسم المنصب", "عدد العلاوات",
    "الرقم الوظيفي", "الرقم الموحد", "الرقم الاحصائي", "طريقة الدفع", "توطين المصرف", "IBAN", "الرفرنس",
    "عدد الاطفال", "حالة الموظف", "الدورة", "Unit_No", "استقطاع مؤقت", "اسم استقطاع المؤقته",
    "اسم اللقب العلمي", "M_Stob", "IdBoss", "حقل مستقبلي تجريبي",
]
COL = {h: i for i, h in enumerate(HEADERS)}

FIRST = ["محمد", "أحمد", "علي", "حسين", "عمر", "خالد", "مصطفى", "حسن", "يوسف", "إبراهيم",
         "سعد", "زياد", "عبد الله", "عبد الرحمن", "مهند", "ليث", "ياسر", "حيدر", "سيف", "رعد"]
FAMILY = ["الدليمي", "العبيدي", "الجميلي", "الكبيسي", "الراوي", "العاني", "الفهداوي", "المحمدي", "النعيمي", "الهيتي"]
RANKS = ["عميد", "عقيد", "مقدم", "رائد", "نقيب", "ملازم أول", "ملازم", "مفوض1", "مفوض3", "مفوض5", "شرطي"]
UNITS = ["قسم النجدة / المقر", "مديرية الشرطة / قسم الإدارة", "فوج الطوارئ الأول", "فوج الطوارئ الثالث",
         "مركز العمليات المشترك", "قسم مكافحة الإجرام / المقر", "مدرسة التدريب", "قسم الأمور الطبية"]
CREDITORS = ["مصرف الرشيد / فرع الرمادي", "مصرف الرشيد / فرع جامعة الانبار", "مديرية تنفيذ الرمادي",
             "مديرية تنفيذ الخالدية", "مديرية تنفيذ الفلوجة", "المصرف العقاري / فرع الانبار",
             "التجهيزات الزراعية / فرع الانبار", "وزارة الداخلية/مكتب الوزير/صندوق شهداء الشرطة"]


def employee_row(name, rank, unit, net, phone_deduction):
    row = [0] * len(HEADERS)
    row[COL["Name_m"]] = name
    row[COL["الرتبة"]] = rank
    row[COL["NO_markaz"]] = unit
    row[COL["حقل مستقبلي تجريبي"]] = "قيمة تجريبية"
    row[COL["NOTA1"]] = ""
    row[COL["اسم المديرية"]] = "مديرية الشرطة"
    row[COL["صافي الراتب"]] = net
    row[COL["اسم الاستقطاع"]] = "شركة الاتصالات" if phone_deduction else " "
    row[COL["مبلغ الاستقطاع"]] = 90000 if phone_deduction else 0
    row[COL["2اسم الاستقطاع"]] = " "
    # عقد واحد بالكاش لاختبار فصل طريقة الدفع، و«عقد» يبقى مختلفاً عن «عقود».
    row[COL["طريقة الدفع"]] = "كاش" if rank == "عقود" else "بطاقة"
    # بنود ورقة الرواتب لتقاريرها (أرقام ثابتة بحسب الاسم حتى تتكرر بين التشغيلات)
    h = sum(ord(c) for c in name)
    row[COL["صندوق الشهداء"]] = 3000 if h % 10 < 8 else 0
    # ٣٥٠٠ لكل فرد من الأسرة؛ موظف واحد بمبلغ لا يقبل القسمة لاختبار التنبيه
    row[COL["الرعاية الصحية"]] = 5000 if name == "عمر خالد جاسم العبيدي" else 3500 * (1 + h % 7) if h % 7 != 6 else 0
    row[COL["نشاط رياضي"]] = 2000 if h % 3 == 0 else 0
    if h % 5 == 0:
        row[COL["2اسم الاستقطاع"]] = "شركة عراق الجوال"
        row[COL["2مبلغ الاستقطاع"]] = 25000 + (h % 7) * 5000
    # المخصصات وبيانات الموظف لتقرير الراتب (أرقام وهمية)
    row[COL["الدرجة الوظيفية"]] = "الدرجة الرابعة"
    row[COL["اسم الوكالة"]] = "وكالة شؤون الشرطة"
    nominal = 400000 + (h % 9) * 15000
    ent = {"الاسمي": nominal, "الخطورة": 485000, "الارزاق": 330000, "الاطفال": 15000 * (h % 4),
           "الزوجية": 70000 if h % 2 else 0, "الشهادة": int(nominal * 0.45), "مؤقتة": 225000}
    for k, v in ent.items():
        row[COL[k]] = v
    row[COL["عدد العلاوات"]] = 1 + h % 3
    row[COL["اسم الشهدة"]] = "بكالوريوس"
    row[COL["اسم استقطاع المؤقته"]] = "مخصصات اضافية/25"
    row[COL["الاستحقاق"]] = sum(ent.values())
    row[COL["توقيفات تقاعدية"]] = nominal // 10
    row[COL["الضريبة"]] = 1283
    row[COL["حماية الاجتماعية"]] = 1273
    deduction_cols = ["توقيفات تقاعدية", "الضريبة", "حماية الاجتماعية", "الرعاية الصحية", "صندوق الشهداء",
                      "نشاط رياضي", "مبلغ الاستقطاع", "2مبلغ الاستقطاع"]
    row[COL["مجموع الاستقطاعات"]] = sum(row[COL[c]] or 0 for c in deduction_cols)
    row[COL["توطين المصرف"]] = "مصرف الرافدين"
    row[COL["IBAN"]] = "IQ00RAFB" + str(10**14 + h * 7919)
    row[COL["الرقم الموحد"]] = str(1000000 + h * 31)
    row[COL["الرقم الوظيفي"]] = str(130000000 + h * 17)
    row[COL["رقم الضمان الصحي"]] = str(400000 + h)
    return row


wb = Workbook(write_only=True)
main = wb.create_sheet("ورقة1")
main.append(HEADERS)

# أسماء وهمية ثابتة للقطات، بينها عقدان للتحقق بصرياً من حصة ١٧٪ في كشف البنك.
# ورقة الاستقطاعات تكتب الاسم الأول دون همزة لاختبار مطابقة الاسم.
fixed = [
    ("أحمد علي حسين الدليمي", "نقيب", "قسم النجدة / المقر", 1398000 if previous else 1452300, True),
    ("عمر خالد جاسم العبيدي", "مفوض3", "فوج الطوارئ الأول", 987650, False),
    ("ليث سعد رعد العاني", "عقود", "قسم النجدة / المقر", 1120000, False),
    ("زياد مهند عمر الراوي", "عقد", "مدرسة التدريب", 1080000, False),
]
names = []
for name, rank, unit, net, phone in fixed:
    main.append(employee_row(name, rank, unit, net, phone))
    names.append(name)

for i in range(count - len(fixed)):
    if i == count // 2:
        main.append(HEADERS)                       # repeated header row, as in exported files
    name = " ".join([rng.choice(FIRST), rng.choice(FIRST), rng.choice(FIRST), rng.choice(FIRST), rng.choice(FAMILY)])
    phone = rng.random() < 0.66
    main.append(employee_row(name, rng.choice(RANKS), rng.choice(UNITS), rng.randrange(700, 3000) * 1000 + rng.randrange(1000), phone))
    names.append(name)

ded = wb.create_sheet("الاستقطاعات")
ded.append(["الرتبة", "الاسم", "اسم الفرع", "المبلغ", "الملاحظات", "القسم"])
ded.append(["نقيب", "احمد علي حسين الدليمي", "مصرف الرشيد / فرع جامعة الانبار", 437024, "قرض شخصي", "قسم النجدة / المقر"])
sheet_deductions = 1
if not previous:
    ded.append(["نقيب", "احمد علي حسين الدليمي", "مديرية تنفيذ الخالدية", 200000, "نفقة 2212 / 2024", "قسم النجدة / المقر"])
    sheet_deductions += 1
for _ in range(count // 4):
    ded.append([rng.choice(RANKS), rng.choice(names[2:]), rng.choice(CREDITORS), rng.randrange(50, 900) * 1000, "", rng.choice(UNITS)])
    sheet_deductions += 1

wb.save(out_path)
with open(expected_path, "w", encoding="utf-8") as f:
    f.write(f"employees={count} deductions={sheet_deductions}\n")
print(f"wrote {out_path}: employees={count} deductions={sheet_deductions}")
